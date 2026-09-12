"""Preserve the official template; persist actual trajectories and full ledgers."""
from __future__ import annotations
import io
import csv
from dataclasses import fields
from datetime import date, datetime, time, timedelta
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from data_io import ATTACHMENT_DIR, coerce_date
from simulation import CommitmentVersion, SimulationConfig, SimulationResult

ARRAY_FIELDS = ("executed", "baseline", "final_commitment", "load_energy", "pv_energy",
                "price", "charge", "discharge", "emergency", "spill", "mode", "soc_before", "soc_after")
DEFAULT_TEMPLATE = ATTACHMENT_DIR / "附件5" / "result3.xlsx"


def _json_value(value):
    if isinstance(value, (date, datetime, Path)):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    raise TypeError(f"cannot serialize {type(value).__name__}")


def _configuration(config):
    return {f.name: getattr(config, f.name) for f in fields(config) if f.name != "data"}


def write_result3(result: SimulationResult, output_path: str | Path,
                  template_path: str | Path = DEFAULT_TEMPLATE) -> Path:
    """Fill only template anchors. Unknown physical values stay blank.

    Adjusted quantities are final commitments; its cost column contains actual
    upward PLUS downward revision fees (baseline is in the preceding sheet).
    Battery blocks use literal calendar left endpoints, never a rotated row.
    Emergency slots list each observed positive interval in a multiline cell,
    with a numeric total; this preserves the fixed illustrative template rows.
    """
    output_path, template_path = Path(output_path), Path(template_path)
    if output_path.resolve() == template_path.resolve():
        raise ValueError("output must not overwrite the attachment template")
    workbook = load_workbook(template_path)
    date_index = {day: row for row, day in enumerate(result.dates)}
    for name, array in (("计划购电量", result.baseline), ("调整购电量", result.final_commitment)):
        sheet = workbook[name]
        for row in range(2, sheet.max_row + 1):
            day = coerce_date(sheet.cell(row, 1).value)
            if day not in date_index:
                continue
            i = date_index[day]
            values = array[i]
            for column, value in enumerate(values, 2):
                sheet.cell(row, column).value = float(value) if np.isfinite(value) else None
            if np.isfinite(values).all():
                sheet.cell(row, 146).value = float(values.sum())
                sheet.cell(row, 147).value = (float(result.price[i] @ values) if name == "计划购电量" else
                    float(sum(v.up_cost + v.down_cost for v in result.versions if v.target_times[0].date() == day)))

    actual = {stamp: (i, j) for i, row in enumerate(result.timestamps)
              for j, stamp in enumerate(row) if result.executed[i, j]}
    sheet = workbook["充放电量"]
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        if not isinstance(value, (date, datetime)):
            continue
        day = coerce_date(value)
        midnight = datetime.combine(day, time())
        for block in range(6):
            times = [midnight + timedelta(hours=4 * block, minutes=10 * k) for k in range(24)]
            if all(stamp in actual for stamp in times):
                for column, name in ((3, "charge"), (4, "discharge")):
                    sheet.cell(row + block, column).value = float(sum(getattr(result, name)[actual[t]] for t in times))
        for offset, stamp in ((0, midnight), (1, midnight + timedelta(days=1))):
            # A state at t is BEFORE execution of [t,t+10m).
            if stamp in actual:
                soc = result.soc_before[actual[stamp]]
            elif stamp - timedelta(minutes=10) in actual:
                soc = result.soc_after[actual[stamp - timedelta(minutes=10)]]
            elif day == result.dates[0] and offset == 0:
                soc = result.config.initial_soc
            else:
                continue
            sheet.cell(row + offset, 6).value = float(soc)

    sheet = workbook["紧急购电量"]
    labels = [workbook["计划购电量"].cell(1, j + 2).value for j in range(144)]
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        if not isinstance(value, (date, datetime)) or coerce_date(value) not in date_index:
            continue
        i = date_index[coerce_date(value)]
        mask = result.executed[i]
        if not mask.any():
            continue
        positives = [j for j in range(144) if mask[j] and result.emergency[i, j] > 1e-7]
        if positives or mask.all():
            prefix = "" if mask.all() else "已执行部分：\n"
            sheet.cell(row, 2).value = prefix + ("\n".join(labels[j] for j in positives) if positives else "无")
            sheet.cell(row, 3).value = float(result.emergency[i, mask].sum())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def save_result(result: SimulationResult, output_dir: str | Path, *, prefix: str = "smoke",
                template_path: str | Path = DEFAULT_TEMPLATE) -> dict[str, Path]:
    """NPZ contains no pickle objects; JSON metadata is embedded for reconstruction."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = dict(solution=output_dir / f"{prefix}_solution.npz", metrics=output_dir / f"{prefix}_metrics.json",
                 trajectory=output_dir / f"{prefix}_trajectory.csv", workbook=output_dir / "result3.xlsx")
    metadata = dict(schema_version=1, dates=[str(d) for d in result.dates], costs=result.costs,
                    config=_configuration(result.config), solve_log=result.solve_log,
                    elapsed_seconds=result.elapsed_seconds,
                    versions=[dict(issued_at=v.issued_at.isoformat(), kind=v.kind,
                                   target_times=[t.isoformat() for t in v.target_times],
                                   baseline_cost=v.baseline_cost, up_cost=v.up_cost, down_cost=v.down_cost)
                              for v in result.versions])
    arrays = {name: getattr(result, name) for name in ARRAY_FIELDS}
    arrays["timestamps"] = np.asarray([[t.isoformat() for t in row] for row in result.timestamps])
    arrays["metadata_json"] = np.asarray(json.dumps(metadata, ensure_ascii=False, default=_json_value, allow_nan=False))
    for key in ("commitment", "revision_up", "revision_down"):
        arrays[f"version_{key}"] = np.asarray([getattr(v, key) for v in result.versions], dtype=float).reshape(-1, 144)
    np.savez_compressed(paths["solution"], **arrays)
    metrics = dict(schema_version=1, dates=metadata["dates"], executed_steps=int(result.executed.sum()),
                   selected_steps=int(result.executed.size), complete=bool(result.executed.all()),
                   costs=result.costs, total_cost=result.total_cost, elapsed_seconds=result.elapsed_seconds,
                   solve_count=len(result.solve_log), optimal_solve_count=sum(bool(v["optimal"]) for v in result.solve_log),
                   config=metadata["config"], final_soc=float(result.soc_after[result.executed][-1]),
                   cost_scope="Full issued commitments and adjacent-version fees; emergency only for executed steps.",
                   calibration="Explicit configuration; no calibration inferred from this result.")
    paths["metrics"].write_text(json.dumps(metrics, indent=2, ensure_ascii=False, default=_json_value, allow_nan=False) + "\n", encoding="utf-8")
    with paths["trajectory"].open("w", newline="", encoding="utf-8-sig") as handle:
        csv_writer = csv.writer(handle)
        csv_writer.writerow(["operating_date", "timestamp", *ARRAY_FIELDS])
        for i, day in enumerate(result.dates):
            for j in range(144):
                values = [getattr(result, name)[i, j] for name in ARRAY_FIELDS]
                csv_writer.writerow([str(day), result.timestamps[i, j].isoformat(),
                                     *["" if isinstance(v, (float, np.floating)) and np.isnan(v) else v for v in values]])
    write_result3(result, paths["workbook"], template_path)
    return paths


def load_result(path: str | Path) -> SimulationResult:
    path = Path(path)
    # Some file-transfer tools prepend a UTF-8 BOM to binary files.
    # Strip only that known wrapper; pickle loading remains disabled.
    raw = path.read_bytes()
    source = raw[3:] if raw.startswith(b"\xef\xbb\xbf") else raw
    with np.load(io.BytesIO(source), allow_pickle=False) as saved:
        metadata = json.loads(str(saved["metadata_json"].item()))
        if metadata["schema_version"] != 1:
            raise ValueError("unsupported solution schema")
        versions = []
        for i, record in enumerate(metadata["versions"]):
            record["issued_at"] = datetime.fromisoformat(record["issued_at"])
            record["target_times"] = tuple(datetime.fromisoformat(t) for t in record["target_times"])
            versions.append(CommitmentVersion(**record, **{name: saved[f"version_{name}"][i].copy()
                for name in ("commitment", "revision_up", "revision_down")}))
        config = metadata["config"]
        config["revision_hours"] = tuple(config["revision_hours"])
        return SimulationResult(dates=tuple(date.fromisoformat(d) for d in metadata["dates"]),
            timestamps=np.asarray([[datetime.fromisoformat(t) for t in row] for row in saved["timestamps"]], dtype=object),
            versions=tuple(versions), config=SimulationConfig(**config), costs=metadata["costs"],
            solve_log=metadata["solve_log"], elapsed_seconds=metadata["elapsed_seconds"],
            **{name: saved[name].copy() for name in ARRAY_FIELDS})
