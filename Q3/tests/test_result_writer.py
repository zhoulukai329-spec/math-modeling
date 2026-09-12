"""Output semantics and independent rejection checks, using real MILP results."""
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta
import importlib
import json
from pathlib import Path
import subprocess
import sys

import numpy as np
from openpyxl import load_workbook
import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
ROOT = SRC.parents[1]
sys.path.insert(0, str(SRC))
from test_simulation import inputs, config
import simulation


def test_output_modules_exist():
    for name in ("make_results", "verify_problem3", "run_problem3"):
        assert (SRC / f"{name}.py").exists(), f"Task 4 {name} is missing"


@pytest.fixture(scope="module")
def result():
    data = inputs(days=4)
    data.load_energy[1:, :] = 4
    return simulation.simulate(config(simulation, data, max_steps=146), data.dates[1], data.dates[2])


@pytest.fixture
def writer():
    return importlib.import_module("make_results")


@pytest.fixture
def verifier():
    return importlib.import_module("verify_problem3")


def test_template_baseline_adjusted_and_unexecuted_are_distinct(writer, result, tmp_path):
    path = writer.write_result3(result, tmp_path / "result3.xlsx")
    workbook = load_workbook(path)
    original = load_workbook(ROOT / "attachment/附件5/result3.xlsx")
    assert workbook.sheetnames == original.sheetnames
    assert workbook.worksheets[0]["B1"].value == "0:10-0:20"
    assert workbook.worksheets[0]["EO1"].value == "0:00-0:10+1"
    assert workbook.worksheets[0]["B3"].value == result.baseline[0, 0]
    assert workbook.worksheets[1]["EO3"].value == result.final_commitment[0, -1]
    assert result.final_commitment[0, -1] != result.baseline[0, -1]
    assert workbook.worksheets[0]["B2"].value is None  # unselected Feb 1
    assert workbook.worksheets[2]["C8"].value is None  # missing Feb 2 00:00
    assert workbook.worksheets[2]["C9"].value == 0  # complete Feb 2 04:00..08:00
    assert workbook.worksheets[2]["F8"].value == result.config.initial_soc
    assert workbook.worksheets[2]["F9"].value == result.soc_before[0, -1]
    assert workbook.worksheets[2]["C14"].value is None  # Mar 20 unexecuted
    assert workbook.worksheets[2]["C21"].value is None  # Dec 31 unexecuted
    assert workbook.worksheets[2]["A20"].value == original.worksheets[2]["A20"].value
    assert workbook.worksheets[0]["B3"].style_id == original.worksheets[0]["B3"].style_id
    assert workbook.worksheets[0]["EQ3"].value == pytest.approx(result.price[0] @ result.baseline[0])
    fees = sum(v.up_cost + v.down_cost for v in result.versions if v.target_times[0].date() == result.dates[0])
    assert workbook.worksheets[1]["EQ3"].value == pytest.approx(fees)


def test_workbook_calendar_blocks_do_not_shift_midnight(writer, result, tmp_path):
    changed = deepcopy(result)
    # Feb 2 04:00 is column 23, and 08:00 (column 47) is excluded.
    changed.charge[0, 23] = 17
    changed.charge[0, 47] = 99
    path = writer.write_result3(changed, tmp_path / "calendar.xlsx")
    assert load_workbook(path).worksheets[2]["C9"].value == 17


def test_npz_roundtrip_and_csv_retain_versions_and_unknowns(writer, verifier, result, tmp_path):
    paths = writer.save_result(result, tmp_path, prefix="smoke")
    restored = writer.load_result(paths["solution"])
    np.testing.assert_equal(restored.executed, result.executed)
    np.testing.assert_allclose(restored.charge, result.charge, equal_nan=True)
    assert restored.versions[-1].issued_at == result.versions[-1].issued_at
    assert restored.config.data is None
    metrics = json.loads(paths["metrics"].read_text(encoding="utf-8"))
    assert metrics["executed_steps"] == 146
    assert metrics["complete"] is False
    assert metrics["total_cost"] == pytest.approx(result.total_cost)
    assert paths["trajectory"].read_text(encoding="utf-8-sig").count("\n") == 289
    assert verifier.verify_solution(restored, workbook_path=paths["workbook"])["passed"]


def test_load_result_accepts_utf8_bom_wrapped_npz(writer, result, tmp_path):
    paths = writer.save_result(result, tmp_path, prefix="bom")
    wrapped = tmp_path / "bom_wrapped.npz"
    wrapped.write_bytes(b"\xef\xbb\xbf" + paths["solution"].read_bytes())
    restored = writer.load_result(wrapped)
    np.testing.assert_equal(restored.executed, result.executed)


@pytest.mark.parametrize("damage,match", [
    ("balance", "balance"), ("soc", "SOC"), ("mode", "mode"),
    ("timestamp", "timestamp"), ("cost", "cost"),
    ("forecast", "forecast"), ("execution_log", "execution"),
    ("unknown", "unexecuted"), ("version", "version"),
    ("past", "frozen"), ("current", "common"),
])
def test_verifier_rejects_independent_tampering(verifier, result, damage, match):
    changed = deepcopy(result)
    if damage == "balance": changed.emergency[0, 0] += 1
    if damage == "soc": changed.soc_before[0, 1] += 1
    if damage == "mode": changed.mode[0, 0] = .5
    if damage == "timestamp": changed.timestamps[0, 0] += timedelta(minutes=10)
    if damage == "cost": changed.costs["baseline"] += 1
    if damage == "forecast": changed.solve_log[0]["forecast_issue"] = "2099-01-01T00:00:00"
    if damage == "execution_log": changed.solve_log.pop()
    if damage == "unknown": changed.charge[1, 20] = 0
    if damage == "version":
        v = changed.versions[1]
        changed.versions = (changed.versions[0], replace(v, revision_up=v.revision_up + 1), *changed.versions[2:])
    if damage == "past":
        v = changed.versions[1]
        commitment = v.commitment.copy(); commitment[0] += 1
        changed.versions = (changed.versions[0], replace(v, commitment=commitment), *changed.versions[2:])
    if damage == "current": changed.solve_log[1]["current_action_disagreement"] = 1
    report = verifier.verify_solution(changed)
    assert report["passed"] is False
    assert any(match.lower() in error.lower() for error in report["errors"]), report


def test_verifier_rejects_workbook_tampering(writer, verifier, result, tmp_path):
    path = writer.write_result3(result, tmp_path / "result3.xlsx")
    workbook = load_workbook(path)
    workbook.worksheets[1]["EO3"] = 999
    workbook.save(path)
    report = verifier.verify_solution(result, workbook_path=path)
    assert not report["passed"]
    assert any("workbook" in error for error in report["errors"])


def test_cli_exposes_modes_and_requires_frozen_full_calibration():
    help_run = subprocess.run([sys.executable, str(SRC / "run_problem3.py"), "--help"], capture_output=True, text=True)
    assert help_run.returncode == 0, help_run.stderr
    assert "experiments" in help_run.stdout and "--max-steps" in help_run.stdout
    full = subprocess.run([sys.executable, str(SRC / "run_problem3.py"), "--mode", "full"], capture_output=True, text=True)
    assert full.returncode != 0
    assert "calibration" in full.stderr.lower()


def test_smoke_rejects_calibration_before_running():
    run = subprocess.run([sys.executable, str(SRC / "run_problem3.py"), "--mode", "smoke",
                          "--calibration", "missing.json", "--max-steps", "1"], capture_output=True, text=True)
    assert run.returncode == 2
    assert "calibration is only" in run.stderr.lower()


@pytest.mark.parametrize("day,row", [("2025-02-01", 2), ("2025-02-02", 8), ("2025-03-20", 14), ("2025-12-31", 21)])
def test_each_battery_template_anchor_uses_actual_calendar_values(writer, result, tmp_path, day, row):
    changed = deepcopy(result)
    delta = datetime.fromisoformat(day).date() - result.dates[0]
    changed.dates = tuple(d + delta for d in result.dates)
    changed.timestamps = np.asarray([[t + delta for t in times] for times in result.timestamps], dtype=object)
    changed.versions = tuple(replace(v, issued_at=v.issued_at + delta,
        target_times=tuple(t + delta for t in v.target_times)) for v in result.versions)
    changed.charge[0, 23] = 17
    path = writer.write_result3(changed, tmp_path / "anchored.xlsx")
    workbook = load_workbook(path)
    assert workbook["充放电量"].cell(row + 1, 3).value == 17
    assert workbook["充放电量"].max_row == 26
    assert workbook["紧急购电量"].max_row == 11


def test_emergency_template_contains_actual_intervals_and_no_added_dates(writer, result, tmp_path):
    path = writer.write_result3(result, tmp_path / "emergency.xlsx")
    sheet = load_workbook(path)["紧急购电量"]
    expected = result.emergency[0, result.executed[0]].sum()
    assert sheet["C5"].value == expected
    assert sheet["B5"].value.startswith("0:10-0:20")
    assert sheet["A8"].value == "⁝"
    assert sheet["A9"].value == datetime(2025, 12, 31)
    assert sheet.max_row == 11


def test_verifier_checks_emergency_binary_mode_even_with_zero_charge(verifier, result):
    changed = deepcopy(result)
    assert changed.emergency[0, 0] > 0 and changed.charge[0, 0] == 0
    changed.mode[0, 0] = 1
    report = verifier.verify_solution(changed)
    assert not report["passed"]
    assert any("emergency" in message and "mode" in message for message in report["errors"])
