from datetime import date, datetime, timedelta
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

import numpy as np
import pytest
from openpyxl import Workbook


SRC_DIR = Path(__file__).resolve().parents[1] / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

import data_io as dio
import forecast as fc


def _input_data(load, pv, hourly_pv=None):
    dates = (date(2025, 1, 1), date(2025, 1, 2), date(2025, 1, 3))
    if hourly_pv is None:
        hourly_pv = {
            (dates[-1], 18 * 60): np.arange(1.0, 25.0),
        }
    return dio.InputData(
        dates=dates,
        load_energy=np.asarray(load, dtype=float),
        pv_energy=np.asarray(pv, dtype=float),
        price=np.ones((len(dates), 144)),
        pv_hourly_forecasts=hourly_pv,
    )


def test_template_left_endpoint_order_is_preserved_without_circular_shift():
    raw_kw = np.arange(144.0).reshape(1, 144)
    energy = dio.power_kw_to_energy(raw_kw)

    assert energy.shape == (1, 144)
    np.testing.assert_allclose(energy[0], raw_kw[0] / 6.0)
    assert dio.template_left_endpoint_minutes()[0] == 10
    assert dio.template_left_endpoint_minutes()[-1] == 24 * 60
    assert dio.template_datetimes(date(2025, 1, 1))[0] == datetime(2025, 1, 1, 0, 10)
    assert dio.template_datetimes(date(2025, 1, 1))[-1] == datetime(2025, 1, 2, 0, 0)


def test_template_order_check_ignores_trailing_total_columns():
    header = ["date"] + [f"{minute // 60}:{minute % 60:02d}" for minute in dio.template_left_endpoint_minutes()]
    header.extend(["daily energy", "daily cost"])

    labels = dio._require_template_order(header)

    assert len(labels) == 144


def test_clock_parser_reads_template_interval_left_endpoint():
    assert dio.parse_clock_minutes("0:10-0:20") == 10
    assert dio.parse_clock_minutes("0:00-0:10+1") == 24 * 60


def test_target_datetime_rolls_18_release_plus_24_hours_to_next_day():
    assert fc.target_datetime(date(2025, 1, 1), "18:00", 24) == datetime(2025, 1, 2, 18, 0)


def test_date_parser_accepts_attachment3_non_padded_dates():
    assert dio.coerce_date("2025-1-1") == date(2025, 1, 1)


def test_information_forecast_interpolates_hourly_pv_without_wrapping_around_day():
    zeros = np.zeros((3, 144))
    data = _input_data(zeros, zeros)

    result = fc.build_information_forecast(
        data, date(2025, 1, 3), "18:00", horizon_steps=4
    )

    assert result.target_times.tolist() == [
        datetime(2025, 1, 3, 18, 0),
        datetime(2025, 1, 3, 18, 10),
        datetime(2025, 1, 3, 18, 20),
        datetime(2025, 1, 3, 18, 30),
    ]
    # Hourly data starts at release + 1h.  The leading value is held, never
    # borrowed from 24h-ahead data by a circular shift; conversion is once.
    np.testing.assert_allclose(result.pv_energy, np.full(4, 1.0 / 6.0))


def test_information_forecast_history_is_cut_off_before_issue_day():
    load = np.zeros((3, 144))
    load[0] = 6.0
    load[1] = 12.0
    load[2] = 9_999.0  # future/current-day actuals must not change the forecast
    pv = np.zeros_like(load)
    data = _input_data(load, pv)

    result = fc.build_information_forecast(
        data, date(2025, 1, 3), "18:00", horizon_steps=2
    )

    # Prior operating days average to 9 kWh; source columns 108 and 109 are
    # the 18:00 and 18:10 left endpoints.  Current-day 9,999 is excluded.
    np.testing.assert_allclose(result.load_energy, [9.0, 9.0])


def test_load_forecast_reuses_q2_week_lag_when_available():
    load = np.stack([np.full(144, float(day + 1)) for day in range(10)])
    dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(10))
    data = dio.InputData(
        dates=dates,
        load_energy=load,
        pv_energy=np.zeros_like(load),
        price=np.ones_like(load),
        pv_hourly_forecasts={(date(2025, 1, 10), 0): np.zeros(24)},
    )

    result = fc.build_information_forecast(
        data, date(2025, 1, 10), "0:00", horizon_steps=2
    )

    # The exact midnight cell belongs to the preceding business row; 00:10
    # belongs to Jan 10's row.  Both use the literal timestamp seven days ago.
    np.testing.assert_allclose(result.load_energy, [2.0, 3.0])


def test_pv_scenarios_use_only_residual_days_before_issue_day():
    load = np.zeros((3, 144))
    pv = np.zeros_like(load)
    data = _input_data(load, pv)
    point = np.array([10.0, 20.0])
    residuals = np.zeros((3, 144))
    residuals[0] = 1.0
    residuals[1] = 2.0
    residuals[2] = 999.0

    scenarios = fc.build_pv_scenarios(
        point, residuals, issue_day_index=2, start_step=0,
        horizon_steps=2, n_scenarios=4, seed=7,
    )

    assert scenarios.shape == (4, 2)
    assert set(np.unique(scenarios - point)).issubset({1.0, 2.0})


def test_load_history_filters_each_cell_by_real_datetime_at_midnight():
    load = np.zeros((3, 144))
    load[0, 143] = 3.0
    load[1, 143] = 9_999.0  # Jan 3 00:00, not available at Jan 3 00:00 issue.
    load[2, 143] = 8_888.0
    data = _input_data(load, np.zeros_like(load), {
        (date(2025, 1, 3), 0): np.zeros(24),
    })

    result = fc.build_information_forecast(data, date(2025, 1, 3), "0:00", horizon_steps=2)

    np.testing.assert_allclose(result.load_energy, [3.0, 0.0])


def test_unique_column_markers_cover_18_2350_next_midnight_and_next_0010():
    markers = np.tile(np.arange(144.0), (3, 1))
    data = _input_data(markers, np.zeros_like(markers))

    result = fc.build_information_forecast(data, date(2025, 1, 3), "18:00", horizon_steps=37)

    np.testing.assert_allclose(result.load_energy[[0, 35, 36]], [107.0, 142.0, 143.0])
    next_step = fc.build_information_forecast(data, date(2025, 1, 3), "18:00", horizon_steps=38)
    assert next_step.target_times[-1] == datetime(2025, 1, 4, 0, 10)
    assert next_step.load_energy[-1] == 0.0


def test_pv_residual_paths_filter_target_midnight_by_real_datetime_and_clip():
    residuals = np.zeros((3, 144))
    residuals[0, 143] = -0.1
    residuals[1, 143] = 999.0  # same calendar instant as issue day 2 midnight

    scenarios = fc.build_pv_scenarios(
        np.array([0.05]), residuals, issue_day_index=2, start_step=143,
        horizon_steps=1, n_scenarios=4, lookback_days=28, seed=3,
    )

    np.testing.assert_allclose(scenarios, 0.0)


@pytest.mark.parametrize("kwargs", [
    {"horizon_steps": 0},
    {"n_scenarios": 0},
    {"lookback_days": 0},
])
def test_pv_scenarios_reject_non_positive_sizes(kwargs):
    args = dict(
        point_pv_energy=np.ones(1), pv_residuals=np.zeros((1, 144)),
        issue_day_index=1, start_step=0, horizon_steps=1,
        n_scenarios=1, lookback_days=1,
    )
    args.update(kwargs)
    if kwargs.get("horizon_steps") == 0:
        args["point_pv_energy"] = np.empty(0)

    with pytest.raises(ValueError, match="positive"):
        fc.build_pv_scenarios(**args)


def _template_header():
    return ["date"] + [f"{minute // 60}:{minute % 60:02d}" for minute in dio.template_left_endpoint_minutes()]


def test_actual_reader_preserves_original_order_and_tail_column():
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        path = Path(directory) / "actual.xlsx"
        book = Workbook()
        ws_load = book.active
        ws_load.title = "load"
        ws_pv = book.create_sheet("pv")
        for ws in (ws_load, ws_pv):
            ws.append(_template_header())
            ws.append([date(2025, 1, 1)] + list(range(144)))
        book.save(path)

        dates, load, pv = dio._read_actuals(path)

    assert dates == (date(2025, 1, 1),)
    np.testing.assert_allclose(load[0], np.arange(144) / 6.0)
    assert load[0, -1] == pytest.approx(143 / 6.0)
    np.testing.assert_allclose(pv, load)


def test_price_reader_rejects_date_mismatch():
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        path = Path(directory) / "price.xlsx"
        book = Workbook()
        ws = book.active
        ws.append(_template_header())
        ws.append([date(2025, 1, 2)] + [1.0] * 144)
        book.save(path)

        with pytest.raises(ValueError, match="price dates do not match"):
            dio._read_price(path, (date(2025, 1, 1),))


def _write_actual_workbook(path, days):
    book = Workbook()
    ws_load = book.active
    ws_load.title = "load"
    ws_pv = book.create_sheet("pv")
    for ws, offset in ((ws_load, 0), (ws_pv, 1_000)):
        ws.append(_template_header())
        for day_index, day in enumerate(days):
            ws.append([day] + list(offset + day_index * 1_000 + np.arange(144)))
    book.save(path)


def _write_price_workbook(path, days):
    book = Workbook()
    ws = book.active
    ws.append(_template_header())
    for day_index, day in enumerate(days):
        ws.append([day] + [day_index + 1.0] * 144)
    book.save(path)


def _write_fixed_price(path):
    book = Workbook()
    book.active.append(["time", "price"])
    for minute in range(10, 1450, 10):
        book.active.append([f"{minute // 60}:{minute % 60:02d}", minute / 1000])
    book.save(path)


def _write_forecast_workbook(path, days, *, invalid_date=None):
    book = Workbook()
    ws = book.active
    ws.append(["date", "release"] + [f"h{i}" for i in range(1, 25)])
    for day_index, day in enumerate(days):
        source_day = invalid_date if invalid_date is not None and day_index == len(days) - 1 else day
        for release_index, release in enumerate(("0:00", "6:00", "12:00", "18:00")):
            ws.append([source_day if release_index == 0 else None, release]
                      + list(100 * day_index + 10 * release_index + np.arange(1, 25)))
    book.save(path)


def test_pv_forecast_reader_preserves_four_releases_and_24_horizons():
    days = (date(2025, 1, 1), date(2025, 1, 2))
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        path = Path(directory) / "forecast.xlsx"
        _write_forecast_workbook(path, days)
        forecasts = dio._read_pv_forecasts(path)

    assert list(forecasts) == [
        (days[0], 0), (days[0], 360), (days[0], 720), (days[0], 1080),
        (days[1], 0), (days[1], 360), (days[1], 720), (days[1], 1080),
    ]
    assert len(forecasts) == 8
    np.testing.assert_allclose(forecasts[(days[1], 1080)], np.arange(131.0, 155.0))


def test_load_inputs_checks_forecast_dates_and_keeps_source_order():
    # Use the repository-local temporary root to avoid a locked system pytest
    # directory on Windows.
    days = (date(2025, 1, 1), date(2025, 1, 2))
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        root = Path(directory)
        (root / "附件5").mkdir()
        _write_actual_workbook(root / "附件2.xlsx", days)
        _write_price_workbook(root / "附件4.xlsx", days)
        _write_fixed_price(root / "附件1.xlsx")
        _write_forecast_workbook(root / "附件3.xlsx", days)
        template = Workbook()
        template.active.append(_template_header() + ["daily energy", "daily cost"])
        template.save(root / "附件5" / "result3.xlsx")

        inputs = dio.load_inputs(root)

        assert inputs.dates == days
        np.testing.assert_allclose(inputs.price, np.tile(np.arange(10, 1450, 10) / 1000, (2, 1)))
        np.testing.assert_allclose(inputs.load_energy[1], (1_000 + np.arange(144)) / 6.0)
        assert inputs.template_labels[0] == "0:10"
        assert inputs.template_labels[-1] == "24:00"
        assert len(inputs.pv_hourly_forecasts) == 8


def test_load_inputs_rejects_forecast_date_not_in_actual_data():
    days = (date(2025, 1, 1), date(2025, 1, 2))
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        root = Path(directory)
        (root / "附件5").mkdir()
        _write_actual_workbook(root / "附件2.xlsx", days)
        _write_price_workbook(root / "附件4.xlsx", days)
        _write_fixed_price(root / "附件1.xlsx")
        _write_forecast_workbook(root / "附件3.xlsx", days, invalid_date=date(2025, 1, 3))
        template = Workbook()
        template.active.append(_template_header() + ["daily energy", "daily cost"])
        template.save(root / "附件5" / "result3.xlsx")

        with pytest.raises(ValueError, match="PV forecast dates do not match"):
            dio.load_inputs(root)


@pytest.mark.parametrize("kwargs", [
    {"operating_dates": (date(2025, 1, 1),)},
    {"issue_time": datetime(2025, 1, 2)},
])
def test_pv_scenarios_reject_unpaired_calendar_arguments(kwargs):
    with pytest.raises(ValueError, match="provided together"):
        fc.build_pv_scenarios(
            np.ones(1), np.zeros((1, 144)), issue_day_index=1, start_step=0,
            horizon_steps=1, n_scenarios=1, lookback_days=1, **kwargs,
        )
