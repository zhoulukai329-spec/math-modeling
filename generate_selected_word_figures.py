"""Regenerate the selected figures that first appeared in the review Word file.

The script reads the completed result workbooks and exports every figure as
PNG (360 dpi), SVG, and PDF. It does not run or alter any optimization model.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


BLUE = "#0072B2"
ORANGE = "#E69F00"
TEAL = "#009E73"
VERMILION = "#D55E00"
PURPLE = "#CC79A7"
BLACK = "#000000"
GRID = "#D9D9D9"


def configure_style() -> None:
    plt.rcParams.update({
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "axes.titlesize": 11,
        "axes.labelsize": 9,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "legend.fontsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "savefig.dpi": 360,
    })


def save_all(fig: plt.Figure, output_dir: Path, stem: str) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for suffix in ("png", "svg", "pdf"):
        fig.savefig(output_dir / f"{stem}.{suffix}", dpi=360,
                    bbox_inches="tight", facecolor="white")
    plt.close(fig)


def read_daily_matrix(path: Path, sheet_name: str) -> tuple[list[datetime], np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    dates: list[datetime] = []
    matrix: list[list[float]] = []
    daily_energy: list[float] = []
    daily_cost: list[float] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        dates.append(row[0])
        matrix.append([float(value or 0.0) for value in row[1:145]])
        daily_energy.append(float(row[-2] or 0.0))
        daily_cost.append(float(row[-1] or 0.0))
    workbook.close()
    return dates, np.asarray(matrix), np.asarray(daily_energy), np.asarray(daily_cost)


def read_emergency(path: Path) -> tuple[list[datetime], np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["紧急购电量"]
    totals: defaultdict[datetime, float] = defaultdict(float)
    current_date = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            current_date = row[0]
        if current_date is not None and row[2] is not None:
            totals[current_date] += float(row[2])
    workbook.close()
    dates = sorted(totals)
    return dates, np.asarray([totals[date] for date in dates])


def find_date_index(dates: list[datetime], month: int = 12, day: int = 21) -> int:
    for index, date in enumerate(dates):
        if date.month == month and date.day == day:
            return index
    raise ValueError(f"result workbook does not contain {month:02d}-{day:02d}")


def month_totals(dates: list[datetime], values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    months = np.arange(2, 13)
    totals = np.asarray([
        values[[date.month == month for date in dates]].sum() / 1000.0
        for month in months
    ])
    return months, totals


def q3_revision_before_after(q3_path: Path, output_dir: Path) -> None:
    dates, baseline, _, _ = read_daily_matrix(q3_path, "计划购电量")
    adjustment_dates, adjustment, adjustment_energy, _ = read_daily_matrix(q3_path, "调整购电量")
    if dates != adjustment_dates:
        raise ValueError("Q3 plan and adjustment dates are not aligned")
    index = find_date_index(dates)
    time = (np.arange(baseline.shape[1]) + 0.5) / 6.0
    final_plan = baseline[index] + adjustment[index]
    fig, ax = plt.subplots(figsize=(7.2, 3.35), layout="constrained")
    ax.step(time, baseline[index] * 6 / 1000, where="mid", color=BLUE,
            lw=1.45, label="0:00基线计划")
    ax.step(time, final_plan * 6 / 1000, where="mid", color=ORANGE,
            lw=1.35, ls="--", label="滚动修订后计划")
    ax.set_title("12月21日滚动修订前后计划购电对比")
    ax.set_xlabel("时刻（h）")
    ax.set_ylabel("购电功率（MW）")
    ax.set_xlim(0, 24)
    ax.set_xticks(np.arange(0, 25, 4))
    ax.grid(axis="y", color=GRID, lw=0.55, alpha=0.7)
    ax.legend(frameon=False, ncol=2, loc="upper left")
    ax.text(0.99, 0.96, f"当日净调整量 {adjustment_energy[index] / 1000:.2f} MWh",
            transform=ax.transAxes, ha="right", va="top", color=ORANGE)
    save_all(fig, output_dir, "q3_fig1_revision_before_after")


def q3_annual_adjustment(q3_path: Path, output_dir: Path) -> None:
    dates, _, daily_adjustment, _ = read_daily_matrix(q3_path, "调整购电量")
    values = daily_adjustment / 1000.0
    maximum = int(np.argmax(values))
    fig, ax = plt.subplots(figsize=(7.2, 3.3), layout="constrained")
    ax.bar(dates, values, width=0.8, color=PURPLE, alpha=0.75)
    ax.set_title("全年滚动修订的每日计划净调整量")
    ax.set_xlabel("日期（2025年）")
    ax.set_ylabel("净调整量（MWh/日）")
    ax.xaxis.set_major_locator(mdates.MonthLocator(interval=2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m月"))
    ax.grid(axis="y", color=GRID, lw=0.55, alpha=0.7)
    ax.annotate(f"峰值 {values[maximum]:.2f} MWh\n{dates[maximum]:%m-%d}",
                xy=(dates[maximum], values[maximum]), xytext=(12, -18),
                textcoords="offset points", color=PURPLE, va="top",
                arrowprops={"arrowstyle": "-", "color": PURPLE, "lw": 0.8})
    save_all(fig, output_dir, "q3_fig2_annual_daily_adjustment")


def q3_monthly_emergency_comparison(q2_path: Path, q3_path: Path, output_dir: Path) -> None:
    q2_dates, q2_emergency = read_emergency(q2_path)
    q3_dates, q3_emergency = read_emergency(q3_path)
    months, q2_monthly = month_totals(q2_dates, q2_emergency)
    _, q3_monthly = month_totals(q3_dates, q3_emergency)
    fig, ax = plt.subplots(figsize=(7.2, 3.35), layout="constrained")
    ax.plot(months, q2_monthly, color="#777777", marker="o", ms=4,
            lw=1.35, label="问题二：日前计划")
    ax.plot(months, q3_monthly, color=TEAL, marker="s", ms=4,
            lw=1.35, ls="--", label="问题三：滚动修订")
    ax.set_title("滚动修订前后各月紧急购电量对比")
    ax.set_xlabel("月份")
    ax.set_ylabel("紧急购电量（MWh/月）")
    ax.set_xticks(months, [f"{month}月" for month in months])
    ax.grid(axis="y", color=GRID, lw=0.55, alpha=0.7)
    ax.legend(frameon=False)
    reduction = 1.0 - q3_emergency.sum() / q2_emergency.sum()
    ax.text(0.99, 0.96, f"全年紧急购电量降低 {reduction:.2%}",
            transform=ax.transAxes, ha="right", va="top", color=TEAL)
    save_all(fig, output_dir, "q3_fig3_monthly_emergency_comparison")


def q4_annual_strategy_comparison(q42_path: Path, q43_path: Path, output_dir: Path) -> None:
    _, _, q42_plan_energy, q42_plan_cost = read_daily_matrix(q42_path, "计划购电量")
    _, _, q43_plan_energy, q43_plan_cost = read_daily_matrix(q43_path, "计划购电量")
    _, _, q43_adjust_energy, q43_adjust_cost = read_daily_matrix(q43_path, "调整购电量")
    _, q42_emergency = read_emergency(q42_path)
    _, q43_emergency = read_emergency(q43_path)
    q42 = np.asarray([q42_plan_cost.sum(), q42_emergency.sum(), q42_plan_energy.sum()])
    q43 = np.asarray([
        q43_plan_cost.sum() + q43_adjust_cost.sum(),
        q43_emergency.sum(),
        q43_plan_energy.sum() + q43_adjust_energy.sum(),
    ])
    relative = q43 / q42 * 100.0
    labels = ["计划及调整费用", "紧急购电量", "计划购电量"]
    x = np.arange(len(labels))
    width = 0.34
    fig, ax = plt.subplots(figsize=(7.2, 3.35), layout="constrained")
    ax.bar(x - width / 2, np.full(3, 100.0), width, color="#777777", label="Q4-2 固定提前期")
    ax.bar(x + width / 2, relative, width, color=TEAL, label="Q4-3 动态提前期")
    for xpos, value in zip(x + width / 2, relative):
        ax.text(xpos, value + 0.35, f"{value:.1f}%", ha="center", va="bottom", fontsize=8)
    ax.set_title("两类动态电价调度策略的全年结果对比")
    ax.set_ylabel("相对Q4-2（%）")
    ax.set_xticks(x, labels)
    ax.set_ylim(80, 103.5)
    ax.grid(axis="y", color=GRID, lw=0.55, alpha=0.7)
    ax.legend(frameon=False, ncol=2)
    save_all(fig, output_dir, "q4_fig1_annual_strategy_comparison")


def q4_representative_day_comparison(q42_path: Path, q43_path: Path, output_dir: Path) -> None:
    q42_dates, q42_plan, _, _ = read_daily_matrix(q42_path, "计划购电量")
    q43_dates, q43_plan, _, _ = read_daily_matrix(q43_path, "计划购电量")
    q43_adjust_dates, q43_adjust, _, _ = read_daily_matrix(q43_path, "调整购电量")
    if q43_dates != q43_adjust_dates:
        raise ValueError("Q4-3 plan and adjustment dates are not aligned")
    q42_index = find_date_index(q42_dates)
    q43_index = find_date_index(q43_dates)
    time = (np.arange(q42_plan.shape[1]) + 0.5) / 6.0
    q43_final = q43_plan[q43_index] + q43_adjust[q43_index]
    fig, ax = plt.subplots(figsize=(7.2, 3.35), layout="constrained")
    ax.step(time, q42_plan[q42_index] * 6 / 1000, where="mid",
            color="#666666", lw=1.4, label="Q4-2 固定提前期计划")
    ax.step(time, q43_final * 6 / 1000, where="mid",
            color=PURPLE, lw=1.35, ls="--", label="Q4-3 动态提前期最终计划")
    ax.set_title("12月21日两类动态电价策略的购电计划对比")
    ax.set_xlabel("时刻（h）")
    ax.set_ylabel("购电功率（MW）")
    ax.set_xlim(0, 24)
    ax.set_xticks(np.arange(0, 25, 4))
    ax.grid(axis="y", color=GRID, lw=0.55, alpha=0.7)
    ax.legend(frameon=False)
    save_all(fig, output_dir, "q4_fig2_representative_day_comparison")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--q2", type=Path, required=True, help="path to result2.xlsx")
    parser.add_argument("--q3", type=Path, required=True, help="path to result3.xlsx")
    parser.add_argument("--q42", type=Path, required=True, help="path to result4-2.xlsx")
    parser.add_argument("--q43", type=Path, required=True, help="path to result4-3.xlsx")
    parser.add_argument("--output-dir", type=Path, default=Path("word_selected_figures"))
    args = parser.parse_args(argv)
    for path in (args.q2, args.q3, args.q42, args.q43):
        if not path.exists():
            raise FileNotFoundError(path)
    configure_style()
    q3_revision_before_after(args.q3, args.output_dir)
    q3_annual_adjustment(args.q3, args.output_dir)
    q3_monthly_emergency_comparison(args.q2, args.q3, args.output_dir)
    q4_annual_strategy_comparison(args.q42, args.q43, args.output_dir)
    q4_representative_day_comparison(args.q42, args.q43, args.output_dir)
    print(f"generated 5 figures in {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
