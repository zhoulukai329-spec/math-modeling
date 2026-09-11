"""Export the four requested dates as paired interval/energy columns."""
import datetime

import numpy as np

import data_io as dio
from make_results import build_emergency_rows


TARGET_DATES = [datetime.date(2025, 3, 20), datetime.date(2025, 6, 21),
                datetime.date(2025, 9, 23), datetime.date(2025, 12, 21)]
DATE_NAMES = [f"{d.year}.{d.month}.{d.day}" for d in TARGET_DATES]
TITLE = "微网在指定日期的紧急购电量"
STEM = dio.OUTPUT_DIR / TITLE


def target_indices(dates):
    base = datetime.date(1899, 12, 30)
    return [int(np.flatnonzero(dates == (day - base).days)[0])
            for day in TARGET_DATES]


def emergency_groups(z):
    groups = []
    for i in target_indices(z["dates"]):
        rows = build_emergency_rows(z["dates"][i:i+1], z["e"][i:i+1])[1:]
        group = [(row[1], row[2]) for row in rows]
        np.testing.assert_allclose(sum(v for _, v in group), z["e"][i].sum(),
                                   rtol=0, atol=1e-4)
        groups.append(group)
    return groups


def wide_rows(groups):
    rows = []
    for j in range(max(1, max(map(len, groups)))):
        row = []
        for group in groups:
            row.extend(group[j] if j < len(group) else ("", None))
        rows.append(row)
    rows.append([v for group in groups for v in ("合计", sum(x[1] for x in group))])
    return rows


def display_rows(groups):
    return [[f"{v:.4f}" if isinstance(v, (float, np.floating)) else (v or "")
             for v in row] for row in wide_rows(groups)]


def markdown_table(groups):
    header = [v for name in DATE_NAMES for v in (name + " 时间段", "购电量 (kWh)")]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|---:|" * 4]
    lines += ["| " + " | ".join(row) + " |" for row in display_rows(groups)]
    return "\n".join(lines)


def export(groups):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    for font in ["Microsoft YaHei", "SimHei", "SimSun"]:
        try:
            font_manager.findfont(font, fallback_to_default=False)
            plt.rcParams["font.sans-serif"] = [font]
            break
        except ValueError:
            continue
    note = ("单位：kWh。连续发生紧急购电的 10 分钟时段合并，购电量为区间内合计；"
            "各日期独立列出，空白表示没有更多区间。数值保留四位小数，合计由未舍入数据计算。")
    STEM.with_suffix(".md").write_text(
        f"# {TITLE}\n\n{note}\n\n{markdown_table(groups)}\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "指定日期紧急购电量"
    ws.append([TITLE])
    ws.merge_cells("A1:H1")
    ws.append([v for name in DATE_NAMES for v in (name, None)])
    for col in [1, 3, 5, 7]:
        ws.merge_cells(start_row=2, end_row=2, start_column=col, end_column=col+1)
    ws.append([v for _ in groups for v in ("时间段", "购电量 (kWh)")])
    raw_rows = wide_rows(groups)
    for row in raw_rows:
        ws.append(row)
    end_row = ws.max_row
    colors = ["E2EEF7", "E2F0E9", "FCEADA", "EEE6F4"]
    border = Border(bottom=Side(style="hair", color="CAD1D8"))
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11, bold=cell.row <= 3 or cell.row == end_row)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if cell.row in [2, 3, end_row]:
                cell.fill = PatternFill("solid", fgColor=colors[(cell.column-1)//2])
            if cell.row >= 4 and cell.column % 2 == 0:
                cell.number_format = "0.0000"
    for j, letter in enumerate("ABCDEFGH"):
        ws.column_dimensions[letter].width = 20 if j % 2 == 0 else 18
    for j in range(1, end_row+1):
        ws.row_dimensions[j].height = 25
    ws.append([note])
    ws.merge_cells(start_row=end_row+1, end_row=end_row+1, start_column=1, end_column=8)
    ws.cell(end_row+1, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[end_row+1].height = 40
    ws.freeze_panes = "A4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:H{end_row+1}"
    wb.save(STEM.with_suffix(".xlsx"))
    saved = load_workbook(STEM.with_suffix(".xlsx"), data_only=True).active
    for r, row in enumerate(raw_rows, start=4):
        for c, value in enumerate(row, start=1):
            actual = saved.cell(r, c).value
            if isinstance(value, (float, np.floating)):
                assert abs(actual-value) < 1e-8
            else:
                assert (actual or "") == (value or "")

    rows = display_rows(groups)
    fig, ax = plt.subplots(figsize=(16, 1.8+0.36*(len(rows)+2)))
    ax.axis("off")
    body = [[v for name in DATE_NAMES for v in (name, "")],
            [v for _ in groups for v in ("时间段", "购电量 (kWh)")]] + rows
    table = ax.table(cellText=body, cellLoc="center", bbox=[0, 0, 1, 0.92])
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("#CAD1D8")
        cell.set_linewidth(0.5)
        if r < 2 or r == len(body)-1:
            cell.set_facecolor("#" + colors[c//2])
            cell.get_text().set_weight("bold")
        elif r % 2 == 0:
            cell.set_facecolor("#F7F8FA")
    ax.set_title(TITLE, fontsize=18, pad=12)
    fig.text(0.5, 0.015, "连续紧急购电时段合并；单位 kWh；合计按未舍入值计算。", ha="center", fontsize=10)
    fig.subplots_adjust(left=0.025, right=0.975, top=0.94, bottom=0.07)
    fig.savefig(STEM.with_suffix(".png"), dpi=180)
    plt.close(fig)


if __name__ == "__main__":
    with np.load(dio.SOLUTION_NPZ) as solution:
        export(emergency_groups(solution))
    print("Specified-date emergency table exported and checked.")
