"""Preserve the official submission template and independently audit all cells."""
from copy import copy
from datetime import datetime, timedelta, time
from pathlib import Path
import os
import tempfile
import numpy as np
from openpyxl import load_workbook
from .calendar_boundary import boundary_errors
from .workbook_contract import preflight_template


def _physical(result):
    rows = {stamp: (float(result.charge[i,j]), float(result.discharge[i,j]),
                    float(result.soc_before[i,j]), float(result.soc_after[i,j]))
            for i,row in enumerate(result.timestamps) for j,stamp in enumerate(row)
            if result.executed[i,j]}
    boundary = getattr(result, "calendar_boundary", None)
    if boundary:
        rows[datetime.fromisoformat(boundary["timestamp"])] = tuple(
            float(boundary[k]) for k in ("charge","discharge","soc_before","soc_after"))
    return rows


def write_workbook(result, output_path, template_path):
    output = Path(output_path)
    preflight_template(template_path, output)
    errors = boundary_errors(result)
    if errors:
        raise ValueError("; ".join(errors))
    book = load_workbook(template_path)
    temporary = None
    try:
        by_day = {day:i for i,day in enumerate(result.dates)}
        adjustment = result.final_commitment - result.baseline
        adjustment[np.abs(adjustment) <= 1e-7] = 0.0
        for name, matrix in (("计划购电量",result.baseline),("调整购电量",adjustment)):
            sheet = book[name]
            for row in range(2,sheet.max_row+1):
                day = sheet.cell(row,1).value.date()
                if day not in by_day:
                    continue
                i = by_day[day]
                for column,value in enumerate(matrix[i],2):
                    sheet.cell(row,column).value = float(value) if np.isfinite(value) else None
                if np.isfinite(matrix[i]).all():
                    sheet.cell(row,146).value = float(matrix[i].sum())
                    sheet.cell(row,147).value = (float(result.price[i] @ matrix[i]) if name == "计划购电量"
                        else float(sum(v.up_cost+v.down_cost for v in result.versions if v.target_times[0].date() == day)))
        physical = _physical(result)
        sheet = book["充放电量"]
        battery_styles = [[copy(sheet.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)]
        battery_heights = [sheet.row_dimensions[r].height for r in range(2,8)]
        sheet.delete_rows(2, sheet.max_row - 1)
        row = 2
        for day in result.dates:
            anchor = datetime.combine(day, time())
            for block in range(6):
                stamps = [anchor+timedelta(hours=4*block,minutes=10*k) for k in range(24)]
                values = [None, f"{4*block}:00-{4*block+4}:00", None, None, None, None]
                if block == 0:
                    values[0] = anchor
                    values[4] = time()
                elif block == 1:
                    values[4] = "24:00"
                for column,index in ((2,0),(3,1)):
                    values[column] = (float(sum(physical[t][index] for t in stamps))
                                      if all(t in physical for t in stamps) else None)
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row + block, column)
                    cell.value = value
                    cell._style = copy(battery_styles[block][column - 1])
                sheet.row_dimensions[row + block].height = battery_heights[block]
            for offset in (0,1):
                t = anchor+timedelta(days=offset)
                value = physical[t][2] if t in physical else (
                    physical[t-timedelta(minutes=10)][3] if t-timedelta(minutes=10) in physical else (
                    result.config.initial_soc if offset == 0 and day == result.dates[0] else None))
                sheet.cell(row+offset,6).value = value
            row += 6
        sheet = book["紧急购电量"]
        emergency_styles = [[copy(sheet.cell(r,c)._style) for c in range(1,4)] for r in range(2,5)]
        emergency_height = sheet.row_dimensions[2].height
        sheet.delete_rows(2, sheet.max_row - 1)
        row = 2
        for i, day in enumerate(result.dates):
            active = result.executed[i]
            indices = np.flatnonzero(active & (result.emergency[i] > 1e-7))
            records = []
            if len(indices):
                for index in indices:
                    index = int(index)
                    start = result.timestamps[i,index]
                    end = start + timedelta(minutes=10)
                    start_prefix = "次日" if start.date() > day else ""
                    end_prefix = "次日" if end.date() > day else ""
                    period = f"{start_prefix}{start.hour}:{start.minute:02d}-{end_prefix}{end.hour}:{end.minute:02d}"
                    records.append((period, float(result.emergency[i,index])))
            else:
                records.append(("无" if active.all() else "已执行部分：无", 0.0))
            for record_index, (period, quantity) in enumerate(records):
                values = [datetime.combine(day,time()) if record_index == 0 else None, period, quantity]
                style_index = 0 if record_index == 0 else min(record_index, 2)
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row, column)
                    cell.value = value
                    cell._style = copy(emergency_styles[style_index][column - 1])
                sheet.row_dimensions[row].height = emergency_height
                row += 1
        output.parent.mkdir(parents=True,exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output.parent,suffix=".xlsx",delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        os.replace(temporary,output)
        temporary = None
        return output
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def verify_workbook(result, path, template_path, require, tolerance):
    """Reconstruct allowed cells separately from the writer; retain all template constants."""
    errors = boundary_errors(result,tolerance)
    for error in errors:
        require(False,error)
    if errors:
        return
    preflight_template(template_path)
    book, source = load_workbook(path), load_workbook(template_path)
    try:
        require(book.sheetnames == source.sheetnames,"workbook sheets")
        if book.sheetnames != source.sheetnames:
            return
        expected = {}
        by_day = {day:i for i,day in enumerate(result.dates)}
        adjustment = result.final_commitment - result.baseline
        adjustment[np.abs(adjustment) <= 1e-7] = 0.0
        for name,matrix in (("计划购电量",result.baseline),("调整购电量",adjustment)):
            for row in range(2,source[name].max_row+1):
                day = source[name].cell(row,1).value.date()
                if day not in by_day:
                    continue
                i = by_day[day]
                for j in range(144):
                    expected[name,row,j+2] = float(matrix[i,j]) if np.isfinite(matrix[i,j]) else None
                if np.isfinite(matrix[i]).all():
                    expected[name,row,146] = float(np.sum(matrix[i]))
                    expected[name,row,147] = (float(np.sum(result.price[i]*matrix[i])) if name == "计划购电量"
                        else float(sum(v.up_cost+v.down_cost for v in result.versions if v.target_times[0].date() == day)))
        # Independent index: do not call the writer's physical lookup.
        timestamps = list(result.timestamps[result.executed])
        values = [list(getattr(result,k)[result.executed]) for k in ("charge","discharge","soc_before","soc_after")]
        boundary = getattr(result,"calendar_boundary",None)
        if boundary:
            timestamps.append(datetime.fromisoformat(boundary["timestamp"]))
            for vector,key in zip(values,("charge","discharge","soc_before","soc_after")):
                vector.append(boundary[key])
        index = {stamp:i for i,stamp in enumerate(timestamps)}
        battery_expected = []
        for day in result.dates:
            anchor = datetime.combine(day,time())
            for block in range(6):
                keys = [anchor+timedelta(minutes=240*block+10*j) for j in range(24)]
                row_values = [anchor if block == 0 else None, f"{4*block}:00-{4*block+4}:00",
                              None, None, time() if block == 0 else "24:00" if block == 1 else None, None]
                row_values[2] = float(np.sum([values[0][index[t]] for t in keys])) if all(t in index for t in keys) else None
                row_values[3] = float(np.sum([values[1][index[t]] for t in keys])) if all(t in index for t in keys) else None
                battery_expected.append((row_values, block))
            for offset in (0,1):
                t = anchor+timedelta(days=offset)
                battery_expected[-6+offset][0][5] = (float(values[2][index[t]]) if t in index else
                    float(values[3][index[t-timedelta(minutes=10)]]) if t-timedelta(minutes=10) in index else
                    result.config.initial_soc if offset == 0 and day == result.dates[0] else None)
        emergency_expected = []
        for i, day in enumerate(result.dates):
            selected = [j for j in range(144) if result.executed[i,j] and result.emergency[i,j] > 1e-7]
            if not selected:
                emergency_expected.append(([datetime.combine(day,time()),
                    "无" if result.executed[i].all() else "已执行部分：无", 0.0], 0))
                continue
            for record_index, j in enumerate(selected):
                start = result.timestamps[i,j]
                end = start + timedelta(minutes=10)
                start_text = ("次日" if start.date() > day else "") + f"{start.hour}:{start.minute:02d}"
                end_text = ("次日" if end.date() > day else "") + f"{end.hour}:{end.minute:02d}"
                emergency_expected.append(([datetime.combine(day,time()) if record_index == 0 else None,
                    f"{start_text}-{end_text}", float(result.emergency[i,j])], min(record_index,2)))
        dynamic = {"充放电量": (battery_expected, 6, 2),
                   "紧急购电量": (emergency_expected, 3, 2)}
        for name,(rows,width,template_start) in dynamic.items():
            sheet, original = book[name], source[name]
            require((sheet.max_row,sheet.max_column) == (1+len(rows),width),f"workbook {name} dimensions")
            require(not sheet.merged_cells,f"workbook {name} merges")
            require(sheet.freeze_panes == original.freeze_panes and sheet.print_area == original.print_area,f"workbook {name} view")
            column_layout = lambda ws: {k:(v.width,v.hidden,v.min,v.max) for k,v in ws.column_dimensions.items()}
            require(column_layout(sheet) == column_layout(original),f"workbook {name} column layout")
            for column in range(1,width+1):
                actual, wanted = sheet.cell(1,column), original.cell(1,column)
                require(actual.value == wanted.value,f"workbook {name}!{actual.coordinate}")
                require(tuple(actual._style or (0,)*9) == tuple(wanted._style or (0,)*9),f"workbook {name}!{actual.coordinate} style")
            for row_number,(row_values,style_index) in enumerate(rows,2):
                template_row = template_start + style_index
                require(sheet.row_dimensions[row_number].height == original.row_dimensions[template_row].height,
                        f"workbook {name} row {row_number} height")
                for column,wanted in enumerate(row_values,1):
                    actual = sheet.cell(row_number,column)
                    valid = (isinstance(actual.value,(int,float)) and np.isfinite(actual.value)
                             and abs(actual.value-wanted) <= tolerance
                             if isinstance(wanted,(int,float,np.number)) else actual.value == wanted)
                    label = f"workbook {name}!{actual.coordinate}"
                    require(valid,label)
                    require(tuple(actual._style or (0,)*9) == tuple(original.cell(template_row,column)._style or (0,)*9),label+" style")
        for original in source:
            if original.title in dynamic:
                continue
            sheet = book[original.title]
            require((sheet.max_row,sheet.max_column) == (original.max_row,original.max_column),f"workbook {sheet.title} dimensions")
            require(str(sheet.merged_cells) == str(original.merged_cells),f"workbook {sheet.title} merges")
            require(sheet.freeze_panes == original.freeze_panes and sheet.print_area == original.print_area,f"workbook {sheet.title} view")
            column_layout = lambda ws: {k:(v.width,v.hidden,v.min,v.max) for k,v in ws.column_dimensions.items()}
            row_layout = lambda ws: {k:(v.height,v.hidden,v.outlineLevel) for k,v in ws.row_dimensions.items()}
            require(column_layout(sheet) == column_layout(original),f"workbook {sheet.title} column layout")
            require(row_layout(sheet) == row_layout(original),f"workbook {sheet.title} row layout")
            for row in original:
                for cell in row:
                    actual = sheet.cell(cell.row,cell.column)
                    wanted = expected.get((sheet.title,cell.row,cell.column),cell.value)
                    valid = (isinstance(actual.value,(int,float)) and np.isfinite(actual.value) and abs(actual.value-wanted) <= tolerance
                             if isinstance(wanted,(int,float,np.number)) else actual.value == wanted)
                    label = f"workbook {sheet.title}!{cell.coordinate}"
                    require(valid,label)
                    require(tuple(actual._style or (0,)*9) == tuple(cell._style or (0,)*9),label+" style")
    finally:
        book.close()
        source.close()
