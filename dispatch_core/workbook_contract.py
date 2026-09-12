"""Preflight the fixed official submission layout before expensive solves."""
from datetime import datetime, timedelta, time
from pathlib import Path
import re
from openpyxl import load_workbook

SHEETS = ("计划购电量", "调整购电量", "充放电量", "紧急购电量")


def preflight_template(template_path, output_path=None):
    template_path = Path(template_path)
    if output_path is not None and Path(output_path).resolve() == template_path.resolve():
        raise ValueError("output must not overwrite the attachment template")
    book = load_workbook(template_path)
    try:
        if tuple(book.sheetnames) != SHEETS:
            raise ValueError("template worksheet names/order do not match the official contract")
        for sheet, shape in zip(book, ((335,147), (335,147), (26,6), (11,3))):
            if (sheet.max_row, sheet.max_column) != shape or sheet.merged_cells:
                raise ValueError(f"template layout mismatch: {sheet.title}")
        for sheet in book.worksheets[:2]:
            if [sheet.cell(1,c).value for c in (1,146,147)] != ["日期\\时间","全天购电量","全天购电费"]:
                raise ValueError("template purchase headers mismatch")
            for i in range(334):
                if sheet.cell(i+2,1).value != datetime(2025,2,1)+timedelta(days=i):
                    raise ValueError(f"template date mismatch: {sheet.title} row {i+2}")
            for j in range(144):
                a, b = (j+1)*10, (j+2)*10
                label = sheet.cell(1,j+2).value
                match = re.fullmatch(r"(\d{1,2}):(\d{1,2})-(\d{1,2}):(\d{1,2})(?:\+1)?", str(label))
                valid = bool(match) and (int(match[1])*60+int(match[2]) == a % 1440
                    and int(match[3])*60+int(match[4]) == b % 1440)
                if not valid:
                    raise ValueError(f"template interval mismatch: {sheet.title} column {j+2}")
        for r, month, day in ((2,2,1),(8,2,2),(14,3,20),(21,12,31)):
            sheet = book["充放电量"]
            if sheet.cell(r,1).value != datetime(2025,month,day):
                raise ValueError("template battery date anchors mismatch")
            if sheet.cell(r,5).value != time() or sheet.cell(r+1,5).value != "24:00":
                raise ValueError("template battery SOC time labels mismatch")
            for k in range(6):
                if sheet.cell(r+k,2).value != f"{4*k}:00-{4*k+4}:00":
                    raise ValueError("template battery blocks mismatch")
        for r, month, day in ((2,2,1),(5,2,2),(9,12,31)):
            if book["紧急购电量"].cell(r,1).value != datetime(2025,month,day):
                raise ValueError("template emergency date anchors mismatch")
        for name, headers in (
            ("充放电量", ["日期","时间段","充电量","放电量","时刻","储电量"]),
            ("紧急购电量", ["日期","购电时间段","购电量"]),
        ):
            if [cell.value for cell in book[name][1]] != headers:
                raise ValueError(f"template headers mismatch: {name}")
        if any(book["充放电量"].cell(20,c).value != "⁝" for c in range(1,7)):
            raise ValueError("template battery separator mismatch")
        if any(book["紧急购电量"].cell(8,c).value != "⁝" for c in range(1,4)):
            raise ValueError("template emergency separator mismatch")
        return template_path
    finally:
        book.close()
