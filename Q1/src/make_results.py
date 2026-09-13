# -*- coding: utf-8 -*-
"""
问题 1 结果生成：
  (a) 填写 result1.xlsx 模板（计划购电量 / 充放电量 两个工作表）
  (b) 输出论文中表 1、表 2 的数值
  (c) 基线对比（无储能）
"""
import sys
from pathlib import Path

#路径
SRC_DIR = Path(__file__).resolve().parent
Q_DIR = SRC_DIR.parent
REPO_ROOT = Q_DIR.parent
ATTACH_DIR = REPO_ROOT / 'attachment'
OUTPUT_DIR = Q_DIR / 'output'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

import os
import zipfile

import numpy as np
from openpyxl import load_workbook

d = np.load(str(OUTPUT_DIR / 'prob1_solution.npz'))
t_min, price = d['t_min'], d['price']
L, G = d['L'], d['G']
x, c, dd, w, E = d['x'], d['c'], d['d'], d['w'], d['E']
cost = float(d['cost'])
N = len(x)

# 表1: 指定时间段购电量
spec = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
        ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]
print('=' * 70)
print('表 1  微网在指定时间段的购电量及全天的购电量和购电费')
print('=' * 70)
for name, k in spec:
    print(f'  {name:12s}  购电量 = {x[k]:10.4f} kWh')
print(f'  {"全天购电量":12s}  = {x.sum():10.4f} kWh')
print(f'  {"全天购电费":12s}  = {cost:10.4f} 元')

# 表2: 指定时间段充放电量
blocks = [('0:00-4:00', 0, 24), ('4:00-8:00', 24, 48), ('8:00-12:00', 48, 72),
          ('12:00-16:00', 72, 96), ('16:00-20:00', 96, 120), ('20:00-24:00', 120, 144)]
print('\n' + '=' * 70)
print('表 2  储能设备在指定时间段的充放电量及 0:00 / 24:00 储电量')
print('=' * 70)
chg_blocks, dis_blocks = [], []
for name, a, b in blocks:
    ch = c[a:b].sum(); di = dd[a:b].sum()
    chg_blocks.append(ch); dis_blocks.append(di)
    print(f'  {name:12s}  充电量 = {ch:10.4f} kWh   放电量 = {di:10.4f} kWh')
print(f'  {"0:00 储电量":12s} = {E[0]:10.4f} kWh')
print(f'  {"24:00 储电量":12s} = {E[-1]:10.4f} kWh')

# 基线：无储能
x_base = np.maximum(0.0, L - G)          # 无储能时每区间购电量 = max(0, 净负载)
w_base = np.maximum(0.0, G - L)          # 无储能时弃光量   = max(0, 光伏富余)
cost_base = float((price * x_base).sum())
saving = cost_base - cost
save_rate = saving / cost_base * 100
print('\n' + '=' * 70)
print('基线对比（无储能）')
print('=' * 70)
print(f'  无储能全天购电量 = {x_base.sum():.4f} kWh,  购电费 = {cost_base:.4f} 元')
print(f'  无储能弃光量     = {w_base.sum():.4f} kWh '
      f'({w_base.sum()/G.sum()*100:.2f}% 光伏), 弃光区间数 = {(w_base > 1e-6).sum()}')
print(f'  有储能全天购电量 = {x.sum():.4f} kWh,  购电费 = {cost:.4f} 元')
print(f'  有储能弃光量     = {w.sum():.6f} kWh')
print(f'  储能节省购电费   = {saving:.4f} 元 (节费率 {save_rate:.2f}%)')
print(f'  弃光量削减       = {w_base.sum() - w.sum():.4f} kWh')
print(f'  能量闭环校验: 弃光削减 - 购电减少 = {(w_base.sum()-w.sum()) - (x_base.sum()-x.sum()):.4f} kWh'
      f'  ≈ 储能净损耗 {c.sum()-dd.sum():.4f} kWh')

TEMPLATE = str(ATTACH_DIR / '附件5' / 'result1.xlsx')
OUT = str(OUTPUT_DIR / 'result1.xlsx')
OUT_TMP = OUT + '.tmp'
OUT_NORMALIZED = OUT + '.normalized.tmp'

try:
    workbook = load_workbook(TEMPLATE)
    plan_sheet = workbook[workbook.sheetnames[0]]
    charge_sheet = workbook[workbook.sheetnames[1]]

    # sheet1: 计划购电量 (B2..B145)
    # 内部数组是自然日顺序：x[0]=0:00-0:10 ... x[143]=23:50-24:00；
    # 官方模板行顺序是：0:10-0:20 ... 23:50-24:00, 次日0:00-0:10。
    # 因此模板第 k+2 行应写 x[(k+1) % N]。
    for k in range(N):
        plan_sheet.cell(row=k + 2, column=2, value=float(x[(k + 1) % N]))

    # sheet2: 充放电量
    for i in range(6):
        charge_sheet.cell(row=i + 2, column=2, value=float(chg_blocks[i]))
        charge_sheet.cell(row=i + 2, column=3, value=float(dis_blocks[i]))
    charge_sheet['E2'] = float(E[0])    # 0:00 储电量
    charge_sheet['E3'] = float(E[-1])   # 24:00 储电量

    workbook.save(OUT_TMP)
    workbook.close()

    # openpyxl 使用包内绝对关系路径；恢复为官方模板的相对路径，兼容 Q1 轻量读取器。
    rels_name = 'xl/_rels/workbook.xml.rels'
    with zipfile.ZipFile(OUT_TMP, 'r') as zin:
        with zipfile.ZipFile(OUT_NORMALIZED, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = zin.read(item.filename)
                if item.filename == rels_name:
                    content = content.replace(
                        b'Target="/xl/worksheets/', b'Target="worksheets/'
                    )
                zout.writestr(item, content)
    os.replace(OUT_NORMALIZED, OUT_TMP)
    os.replace(OUT_TMP, OUT)
finally:
    if os.path.exists(OUT_TMP):
        os.remove(OUT_TMP)
    if os.path.exists(OUT_NORMALIZED):
        os.remove(OUT_NORMALIZED)
print(f'\n已生成结果文件: {OUT}')
print('工作表“计划购电量”: 144 个 10 分钟区间购电量 (B2:B145)')
print('工作表“充放电量”:   6 个 4 小时块充/放电量 (B2:C7) + 0:00/24:00 储电量 (E2:E3)')
