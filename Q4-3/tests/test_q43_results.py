from __future__ import annotations

from datetime import date, timedelta
import importlib
from pathlib import Path
import sys

import numpy as np
from openpyxl import load_workbook

SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(SRC))


def _load():
    if str(SRC) in sys.path:
        sys.path.remove(str(SRC))
    sys.path.insert(0, str(SRC))
    for name in ("data_io", "forecast", "price_forecast", "simulation", "make_results", "verify_problem43"):
        sys.modules.pop(name, None)
    dio = importlib.import_module("data_io")
    sim = importlib.import_module("simulation")
    writer = importlib.import_module("make_results")
    verifier = importlib.import_module("verify_problem43")
    return dio, sim, writer, verifier


def _result(dio, sim):
    days = 3
    dates = tuple(date(2025, 2, 1) + timedelta(days=i) for i in range(days))
    load = np.full((days, 144), 2.0)
    price = np.full((days, 144), 1.25)
    data = dio.InputData(dates, load, np.zeros_like(load), price, {
        (day, minute): np.zeros(24) for day in dates for minute in (0, 360, 720, 1080)
    })
    cfg = sim.SimulationConfig(data=data, horizon_steps=3, max_steps=3, n_scenarios=1,
        deterministic=True, january_warmup=False, backend="event-policy",
        initial_soc=0, soc_min=0, soc_max=10, charge_limit=0, discharge_limit=0,
        terminal_soc=0, terminal_penalty=0)
    return sim.simulate(cfg, dates[1], dates[1])


def test_q43_writer_uses_result43_names_and_roundtrips(tmp_path):
    dio, sim, writer, verifier = _load()
    result = _result(dio, sim)
    paths = writer.save_result43(result, tmp_path, prefix="smoke")
    assert paths["workbook"].name == "result4-3.xlsx"
    restored = writer.load_result(paths["solution"])
    np.testing.assert_allclose(restored.price_forecast, result.price_forecast, equal_nan=True)
    report = verifier.verify_solution43(restored, workbook_path=paths["workbook"],
                                        template_path=writer.DEFAULT_TEMPLATE)
    assert report["passed"], report["errors"]


def test_q43_verifier_recomputes_dynamic_actual_price_cost():
    dio, sim, _, verifier = _load()
    result = _result(dio, sim)
    result.costs["baseline"] += 1.0
    report = verifier.verify_solution43(result)
    assert not report["passed"]
    assert any("cash cost baseline" in error for error in report["errors"])


def test_q43_emergency_uses_separate_ten_minute_rows(tmp_path):
    dio, sim, writer, _ = _load()
    result = _result(dio, sim)
    result.emergency[0, 0] = 1
    result.emergency[0, 1] = 2
    path = writer.write_result43(result, tmp_path / "result4-3.xlsx")
    workbook = load_workbook(path, data_only=True)
    try:
        emergency = workbook["紧急购电量"]
        assert emergency.max_row == 3
        assert emergency["B2"].value == "0:10-0:20"
        assert emergency["C2"].value == 1
        assert emergency["B3"].value == "0:20-0:30"
        assert emergency["C3"].value == 2
    finally:
        workbook.close()
